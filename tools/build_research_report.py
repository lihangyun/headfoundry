from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "report-source.md"
LEDGER = ROOT / "docs" / "claim-source-ledger.md"
OUTPUT = ROOT / "docs" / "HeadFoundry_Technical_Research.docx"

BLUE = "1F4E78"
DARK_BLUE = "17365D"
LIGHT_BLUE = "DDEBF7"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "667085"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, width=9360, indent=120) -> None:
    properties = table._tbl.tblPr
    table_width = properties.first_child_found_in("w:tblW")
    table_width.set(qn("w:w"), str(width))
    table_width.set(qn("w:type"), "dxa")
    table_indent = OxmlElement("w:tblInd")
    table_indent.set(qn("w:w"), str(indent))
    table_indent.set(qn("w:type"), "dxa")
    properties.append(table_indent)
    for row in table.rows:
        for cell in row.cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_field(run, instruction: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, text, end])


def add_hyperlink(paragraph, label: str, url: str) -> None:
    relationship = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.extend([color, underline])
    run.append(props)
    node = OxmlElement("w:t")
    node.text = label
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, text: str) -> None:
    pattern = re.compile(r"\[([^]]+)]\((https?://[^)]+)\)|`([^`]+)`|\*\*([^*]+)\*\*")
    cursor = 0
    for match in pattern.finditer(text):
        paragraph.add_run(text[cursor : match.start()])
        if match.group(1):
            add_hyperlink(paragraph, match.group(1), match.group(2))
        elif match.group(3):
            run = paragraph.add_run(match.group(3))
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        else:
            paragraph.add_run(match.group(4)).bold = True
        cursor = match.end()
    paragraph.add_run(text[cursor:])


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.1
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    if "Masthead" not in document.styles:
        masthead = document.styles.add_style("Masthead", WD_STYLE_TYPE.PARAGRAPH)
        masthead.font.name = "Calibri"
        masthead.font.size = Pt(9)
        masthead.font.bold = True
        masthead.font.color.rgb = RGBColor.from_string(WHITE)
        masthead.paragraph_format.space_after = Pt(0)


def add_header_footer(section) -> None:
    for footer in (section.footer, section.even_page_footer, section.first_page_footer):
        footer.is_linked_to_previous = False
        paragraph = footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run("Confidential working brief  •  ")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MID_GRAY)
        page_run = paragraph.add_run()
        add_field(page_run, "PAGE")


def add_masthead(document: Document) -> None:
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(4.7)
    table.columns[1].width = Inches(1.7)
    left, right = table.rows[0].cells
    set_cell_shading(left, BLUE)
    set_cell_shading(right, DARK_BLUE)
    set_cell_margins(left, 130, 170, 130, 170)
    set_cell_margins(right, 130, 170, 130, 170)
    left_p = left.paragraphs[0]
    left_p.style = "Masthead"
    left_p.add_run("HEADFOUNDRY  /  DECISION BRIEF")
    right_p = right.paragraphs[0]
    right_p.style = "Masthead"
    right_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right_p.add_run("03 SEP 2026")


def add_summary_table(document: Document) -> None:
    table = document.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    set_table_width(table)
    entries = (
        ("Decision", "Proceed with a clean-room, optimization-first reconstruction system."),
        ("Quality claim", "UNVERIFIED until clay, texture, protected views, mesh health, and full reconstruction all pass."),
        ("First gate", "Synthetic camera projection must remain below 0.05 px p95 reprojection error."),
        ("Primary risk", "Commercially cleared full-head data and dense correspondence training assets."),
    )
    for row, (label, value) in zip(table.rows, entries, strict=True):
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        row.cells[0].width = Inches(1.35)
        row.cells[1].width = Inches(5.05)
        row.cells[0].paragraphs[0].add_run(label).bold = True
        row.cells[1].paragraphs[0].add_run(value)


def add_markdown_body(document: Document, text: str) -> None:
    lines = text.splitlines()
    for index, raw in enumerate(lines):
        line = raw.strip()
        if not line or line.startswith("Date:") or line.startswith("Decision:") or line.startswith("Quality status:"):
            continue
        if line.startswith("# "):
            continue
        if line == "## Sources":
            document.add_heading("Sources and audit trail", level=1)
            continue
        if line.startswith("## "):
            document.add_heading(line[3:], level=1)
            continue
        if line.startswith("### "):
            document.add_heading(line[4:], level=2)
            continue
        numbered = re.match(r"^(\d+)\.\s+(.*)$", line)
        if numbered:
            paragraph = document.add_paragraph(style="List Number")
            add_inline(paragraph, numbered.group(2))
            continue
        if line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, line[2:])
            continue
        paragraph = document.add_paragraph()
        add_inline(paragraph, line)


def parse_ledger_rows(text: str) -> list[list[str]]:
    rows = []
    for line in text.splitlines():
        if not line.startswith("|"):
            continue
        cells = [cell.strip() for cell in line.strip("|").split("|")]
        if cells[0] == "Claim" or set(cells[0]) == {"-"}:
            continue
        rows.append(cells)
    return rows


def add_ledger(document: Document) -> None:
    document.add_page_break()
    document.add_heading("Claim-to-source ledger", level=1)
    paragraph = document.add_paragraph("Primary evidence used for the architecture decision. Accessed 03 September 2026.")
    paragraph.runs[0].italic = True
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ("Claim", "Primary source", "Evidence note")
    for cell, label in zip(table.rows[0].cells, headers, strict=True):
        set_cell_shading(cell, LIGHT_GRAY)
        cell.paragraphs[0].add_run(label).bold = True
    for claim, source, note in parse_ledger_rows(LEDGER.read_text(encoding="utf-8")):
        cells = table.add_row().cells
        add_inline(cells[0].paragraphs[0], claim)
        add_inline(cells[1].paragraphs[0], source)
        add_inline(cells[2].paragraphs[0], note)
    set_table_width(table)
    for row in table.rows:
        row.cells[0].width = Inches(2.75)
        row.cells[1].width = Inches(1.75)
        row.cells[2].width = Inches(1.9)
        for cell in row.cells:
            set_cell_margins(cell, top=35, start=90, bottom=35, end=90)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.size = Pt(7.5)


def build() -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    configure_styles(document)
    add_header_footer(section)
    add_masthead(document)

    title = document.add_paragraph()
    title.paragraph_format.space_before = Pt(20)
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("KeenTools-class multi-view 3D head reconstruction")
    run.font.name = "Calibri"
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    subtitle = document.add_paragraph("Technical research, legal boundary, architecture, and falsifiable delivery plan")
    subtitle.paragraph_format.space_after = Pt(16)
    subtitle.runs[0].font.size = Pt(12)
    subtitle.runs[0].font.color.rgb = RGBColor.from_string(MID_GRAY)
    add_summary_table(document)

    add_markdown_body(document, SOURCE.read_text(encoding="utf-8"))
    add_ledger(document)

    document.core_properties.title = "HeadFoundry technical research"
    document.core_properties.subject = "Clean-room multi-view 3D head reconstruction"
    document.core_properties.author = "HeadFoundry"
    document.core_properties.keywords = "3D head reconstruction, camera estimation, texture fusion, fixed topology"
    document.save(OUTPUT)


if __name__ == "__main__":
    build()
